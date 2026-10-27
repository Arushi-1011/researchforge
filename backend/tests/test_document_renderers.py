"""Unit tests for document renderers without a database or external APIs."""
from io import BytesIO
from docx import Document
from app.modules.documents.renderers import render_docx, render_html


def sample_content():
    return {
        "title": "ResearchForge <Report>",
        "created_at": "2026-10-27 00:00 UTC",
        "project_name": "Demo",
        "sections": [
            {"key": "overview", "title": "1. Project Overview", "kind": "list",
             "items": ["Project name: Demo"], "text": "A <script>alert(1)</script>"},
            {"key": "literature", "title": "2. Literature Table", "kind": "table",
             "headers": ["#", "Title"], "rows": [["1", "A Paper"]], "items": [], "text": ""},
            {"key": "gaps", "title": "5. Research Gaps", "kind": "list",
             "items": [], "text": ""},
        ],
    }


def test_html_renderer_escapes_untrusted_content():
    html = render_html(sample_content())
    assert "&lt;script&gt;" in html
    assert "<script>alert(1)</script>" not in html
    assert "A Paper" in html


def test_docx_renderer_contains_title_and_section_content():
    payload = render_docx(sample_content())
    document = Document(BytesIO(payload))
    paragraphs = [p.text for p in document.paragraphs]
    assert "ResearchForge <Report>" in paragraphs
    assert "1. Project Overview" in paragraphs
    assert "5. Research Gaps" in paragraphs
    assert document.tables[0].rows[1].cells[1].text == "A Paper"
